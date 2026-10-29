"""Independent desert tests with explicit observation support and retained failures.

Run after v03 reproduction. Inputs are checked against the primary PMC manifest.
See reviews/v05/external-evaluation-protocol.md for the retrospective design.
"""
from pathlib import Path
from contextlib import closing
import hashlib, json, sys
from dataclasses import replace
from itertools import product
import numpy as np
import pandas as pd
from scipy.optimize import minimize_scalar
import openpyxl
from eq_biodt.v03.process import Parameters, gas_exchange
ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'reviews/v05'; DATA=ROOT/'data/external'
P=Parameters(k_h2_s=.001,k_co_s=0.)
SEED=20261029
CANDIDATES=[replace(P,porosity=phi,retention_b=b,psi_entry_mpa=entry,q10=q)
            for phi,b,entry,q in product([.35,.4,.5],[2.,4.,6.],[-.0001,-.001,-.01],[1.,2.])]

def verified(study,name,data_dir=DATA):
 records=json.loads((data_dir/'pmc-supplement-manifest.json').read_text())
 record=next((r for r in records if r['study']==study and r['file']==name),None)
 if record is None:
  raise ValueError(f"Unknown external source: {study}/{name}")
 path=data_dir/study/name
 body=path.read_bytes()
 if len(body)!=record['bytes'] or hashlib.sha256(body).hexdigest()!=record['sha256'] or hashlib.md5(body).hexdigest()!=record['md5']:
  raise ValueError(f"External source digest mismatch: {path}")
 return path

def verify_source_record(key,sources_dir=ROOT/'data/sources'):
 info=json.loads((sources_dir/'manifest.json').read_text())['sources'][key]
 path=sources_dir/info['file']
 if hashlib.sha256(path.read_bytes()).hexdigest()!=info['sha256']:
  raise ValueError(f"Source digest mismatch: {key}")
 return path

def flux(theta,temp,p=P):
 return gas_exchange(float(theta),float(temp),p)['h2_mol_m2_s']*1e9

def mean_ci(values,rng):
 v=np.asarray(values,float)
 draws=rng.choice(v,(5000,len(v)),replace=True).mean(axis=1)
 return float(v.mean()),*map(float,np.quantile(draws,[.025,.975]))

def evaluate_sonoran_fold(train,test):
 if train.empty or test.empty or not set(train.campaign).isdisjoint(set(test.campaign)):
  raise ValueError("Train and test campaigns must be non-empty and disjoint")
 def predict(g,p):return np.array([flux(r.theta,r.temp,p) for r in g.itertuples()])
 def fit(g,p):
  def loss(logk):return np.mean((predict(g,replace(p,k_h2_s=10**logk))-g.y.to_numpy())**2)
  opt=minimize_scalar(loss,bounds=(-9,2),method='bounded',options={'xatol':1e-9})
  if not opt.success:
   raise RuntimeError("Scalar kinetic fit failed to converge")
  return replace(p,k_h2_s=10**opt.x),float(opt.fun)
 calibrated,_=fit(train,P)
 diff,_=fit(train,replace(P,mechanism='diffusion_only'))
 tuned,loss=min((fit(train,p) for p in CANDIDATES),key=lambda x:x[1])
 design=lambda g: np.column_stack([np.ones(len(g)),g.theta,g.temp])
 beta=np.linalg.lstsq(design(train),train.y,rcond=None)[0]
 preds={'frozen_reference':predict(test,P),'calibrated_kinetics':predict(test,calibrated),
        'calibrated_diffusion_only':predict(test,diff),'calibrated_retention':predict(test,tuned),
        'training_mean':np.repeat(train.y.mean(),len(test)),'linear_moisture_temperature':design(test)@beta}
 fit_record={'heldout_campaign':str(test.campaign.iloc[0]),'training_campaigns':sorted(train.campaign.unique()),
             'kinetics_k':calibrated.k_h2_s,'diffusion_k':diff.k_h2_s,
             'tuned':{'k':tuned.k_h2_s,'phi':tuned.porosity,'b':tuned.retention_b,'entry':tuned.psi_entry_mpa,'q10':tuned.q10},
             'tuned_training_mse':loss,'linear_coefficients':beta.tolist()}
 return preds,fit_record

def sonoran():
 rng=np.random.default_rng(SEED)
 lab=pd.read_csv(verified('buzzard2022','peerj-10-12966-s007.csv'))
 field=pd.read_csv(verified('buzzard2022','peerj-10-12966-s008.csv'))
 # Authors' supplementary code identifies Hydrogen.Uptake as updated calibration.
 for df in (lab,field):df['uptake_nmol_m2_s']=-df['Hydrogen.Uptake']
 lab['gwc']=lab['Moisture.level'].map({'Dry (0.02 GWC)':.02,'Moist (0.10 GWC)':.10,'Wet (0.2 GWC)':.20})
 if not (lab.gwc.notna().all() and len(lab)==48 and len(field)==64 and np.isfinite(lab.uptake_nmol_m2_s).all() and np.isfinite(field.uptake_nmol_m2_s).all()):
  raise ValueError("Invalid Sonoran incubation or field observations")
 lab['source_batch']=lab['Treatment']+':'+lab['Batch'].astype(str)
 rows=[]
 for treatment,g in lab.groupby('Treatment'):
  pivot=g.pivot(index='Batch',columns='gwc',values='uptake_nmol_m2_s')
  if len(pivot)!=4 or pivot.isna().any().any():
   raise ValueError(f"Incomplete Sonoran source-batch design for {treatment}")
  for lo,hi in [(0.02,.10),(.10,.20),(.02,.20)]:
   mean,lcl,ucl=mean_ci(pivot[hi]-pivot[lo],rng)
   rows.append({'study':'Sonoran urban','site':treatment,'dry_gwc':lo,'wet_gwc':hi,
                'n_source_batches':len(pivot),'difference':mean,'lcl':lcl,'ucl':ucl,
                'unit':'nmol H2 m-2 s-1','scope':'Within-plot source-batch bootstrap; one residence'})
 contrasts=pd.DataFrame(rows)
 contrasts.to_csv(OUT/'sonoran_hydration_contrasts.tsv',sep='\t',index=False)
 # GWC-to-VWC conversion needs bulk density: explicitly assumed, not measured.
 sensitivity=[]
 for rho,phi,b,entry in product([1.2,1.5,1.8],[.4,.5,.6],[2.,4.,6.],[-.0001,-.001,-.01]):
  p=replace(P,porosity=phi,retention_b=b,psi_entry_mpa=entry)
  for lo,hi in [(0.02,.10),(.10,.20),(.02,.20)]:
   a,z=flux(lo*rho,25,p),flux(hi*rho,25,p)
   sensitivity.append({'bulk_density_g_cm3':rho,'porosity':phi,'retention_b':b,'entry_mpa':entry,
                       'dry_gwc':lo,'wet_gwc':hi,'predicted_difference':z-a,
                       'dry_pred':a,'wet_pred':z})
 pd.DataFrame(sensitivity).to_csv(OUT/'sonoran_input_sensitivity.tsv',sep='\t',index=False)
 # Same sensor value is repeated for four collars: split all rows of each campaign.
 field['campaign']=pd.to_datetime(field.Date).dt.strftime('%Y-%m-%d')
 if not field.groupby(['campaign','Treatment'])[['sm','temp']].nunique().eq(1).all().all():
  raise ValueError("Sensor moisture/temperature must be constant across collars within each plot-campaign")
 means=field.groupby(['campaign','Treatment'],as_index=False).agg(theta=('sm','first'),temp=('temp','first'),y=('uptake_nmol_m2_s','mean'),n=('Batch','size'))
 if len(means)!=16 or not means.n.eq(4).all():
  raise ValueError("Expected 16 plot-campaign means of 4 collars each")
 predictions=[]; fits=[]
 for campaign in sorted(means.campaign.unique()):
  train=means[means.campaign.ne(campaign)];test=means[means.campaign.eq(campaign)]
  preds,fit_record=evaluate_sonoran_fold(train,test)
  for model,pred in preds.items():
   for row,value in zip(test.itertuples(),pred):
    predictions.append({'campaign':campaign,'plot':row.Treatment,'model':model,'observed':row.y,'predicted':float(value)})
  fits.append(fit_record)
 pred=pd.DataFrame(predictions);pred.to_csv(OUT/'sonoran_field_predictions.tsv',sep='\t',index=False)
 metrics=[]
 for model,g in pred.groupby('model'):
  metrics.append({'model':model,'rmse':float(np.sqrt(np.mean((g.observed-g.predicted)**2))),'mae':float(np.mean(abs(g.observed-g.predicted)))})
 lab.to_csv(OUT/'sonoran_lab_observations.tsv',sep='\t',index=False)
 means.to_csv(OUT/'sonoran_field_observations.tsv',sep='\t',index=False)
 return {'lab_bottles':48,'field_collars':64,'field_plot_campaign_means':16,'independent_residences':1,
         'outcome_column':'Hydrogen.Uptake, updated response-factor calibration; uptake positive after sign reversal',
         'field_validation':'Leave-one-campaign-out, four entire campaigns; descriptive errors on 16 plot-campaign means',
         'field_metrics':metrics,'fits':fits,'lab_default_assumed_density_g_cm3':1.5,
         'lab_default_predictions':{str(gwc):flux(gwc*1.5,25) for gwc in [.02,.10,.20]},
         'lab_scope':'Controlled hydration; same source batches at three water contents. Uncertainty is within plots, not among replicated treatments or desert sites.'}

def namib():
 path=verified('tribbia2026','aem.00265-26-s0002.xlsx')
 with closing(openpyxl.load_workbook(path,read_only=True,data_only=True)) as w:
  rows=[]
  for state,sheet in [('dry','Table S7. Dry H2 oxidation'),('wet','Table S8. Wet H2 oxidation')]:
   counts={}
   for vals in w[sheet].iter_rows(min_row=6,values_only=True):
    site,zone,k,rate,biomass,*_=vals
    if not isinstance(site,str) or not site.startswith('C14-'):continue
    counts[site]=counts.get(site,0)+1
    rows.append({'site':site,'zone':zone,'state':state,'replicate':counts[site],'k_per_hour':k,'nmol_h_g':rate,'nmol_h_16s_copy':biomass})
  moisture={}
  for vals in w['Table S1. Gravimetric SMC%'].iter_rows(min_row=5,values_only=True):
   if isinstance(vals[0],str) and vals[0].startswith('C14-') and vals[3] is not None:moisture[vals[0]]=float(vals[3])
 rates=pd.DataFrame(rows)
 if len(rates)!=60 or len(moisture)!=10:
  raise ValueError("Unexpected Namib rate or moisture row counts")
 # Cross-check the declared gas-volume and mass observation map.
 expected=rates.k_per_hour*.0047436*530e-9*1e9/2
 # Header gas amount is rounded: allow 5 ppm relative discrepancy.
 if not np.allclose(expected,rates.nmol_h_g,rtol=5e-6,atol=0):
  raise ValueError("Namib rate conversion does not match declared headspace/mass formula")
 rates.to_csv(OUT/'namib_observations.tsv',sep='\t',index=False)
 rng=np.random.default_rng(SEED);contrasts=[]
 for site in sorted(rates.site.unique(),key=lambda x:int(x.split('-')[1])):
  dry=rates[(rates.site==site)&rates.state.eq('dry')].nmol_h_g.to_numpy()
  wet=rates[(rates.site==site)&rates.state.eq('wet')].nmol_h_g.to_numpy()
  if not len(dry):
   contrasts.append({'site':site,'n_dry_rates':0,'n_wet_rates':len(wet),'wet_mean':float(wet.mean()),'status':'Dry sub-atmospheric endpoint not attained within 1500 h; no numerical zero imputed'})
   continue
  # Dry and wet assays used distinct subsamples: independent within-site draws.
  draws=rng.choice(wet,(5000,len(wet))).mean(axis=1)/rng.choice(dry,(5000,len(dry))).mean(axis=1)
  contrasts.append({'site':site,'n_dry_rates':len(dry),'n_wet_rates':len(wet),'dry_mean':float(dry.mean()),'wet_mean':float(wet.mean()),'fold':float(wet.mean()/dry.mean()),'fold_lcl':float(np.quantile(draws,.025)),'fold_ucl':float(np.quantile(draws,.975)),'status':'Independent wet/dry subsamples; within-site bootstrap'})
 pd.DataFrame(contrasts).to_csv(OUT/'namib_hydration_contrasts.tsv',sep='\t',index=False)
 # Gravimetric content in workbook cells is a fraction with Excel percent format.
 predictions=[]
 for site,gwc in moisture.items():
  # Additional 0.4 mL to 2g native soil; convert to dry-soil denominator.
  wetted_gwc=gwc+.2*(1+gwc)
  for rho in [1.2,1.5,1.8]:
   dry=flux(gwc*rho,30);wet=flux(wetted_gwc*rho,30)
   predictions.append({'site':site,'dry_gwc':gwc,'wet_gwc':wetted_gwc,'assumed_bulk_density_g_cm3':rho,'dry_reference_flux':dry,'wet_reference_flux':wet,'predicted_direction':int(np.sign(wet-dry)),
                       'scope':'Conditional field-layer response direction only; bottle rate units not equated with areal flux'})
 pd.DataFrame(predictions).to_csv(OUT/'namib_conditional_predictions.tsv',sep='\t',index=False)
 return {'sites':10,'dry_sites_with_quantified_atmospheric_rates':2,'wet_sites_with_quantified_atmospheric_rates':10,'rate_rows':len(rates),'contrasts':contrasts,
         'prediction_scope':'Hydration direction under declared density/retention assumptions. Areal model values are not scored against bottle units.',
         'measurement_exclusion':'Historical site water potentials and climatological grid moisture do not measure incubation conditions; unavailable bottle potentials prevent a direct -85 MPa threshold test.',
         'reference_failure':'Default retention and density conversion predict zero dry activity at sites where positive dry atmospheric oxidation was measured; this tests the assumed conversion plus activation, not a measured water-potential threshold alone.'}

def evaluate_temperate_folds(d):
 rows=[]
 for soil in sorted(d.soil.unique()):
  train=d[d.soil.ne(soil)];test=d[d.soil.eq(soil)]
  if train.empty or test.empty or not set(train.soil).isdisjoint(set(test.soil)):
   raise ValueError("Train and test soils must be non-empty and disjoint")
  for model in ['diffusion_only','activation_diffusion','training_mean']:
   if model=='training_mean':amplitude=None;pred=np.repeat(train.bottle_k_min.mean(),len(test))
   else:
    x=train[model].to_numpy();amplitude=max(0.,float(x@train.bottle_k_min/(x@x)));pred=test[model]*amplitude
   for r,value in zip(test.itertuples(),pred):rows.append({'heldout_soil':soil,'sample_id':r.sample_id,'replicate':r.replicate,'model':model,'amplitude':amplitude,'observed_k_min':r.bottle_k_min,'predicted_k_min':float(value)})
 return pd.DataFrame(rows)

def temperate():
 verify_source_record('Reji_bottles')
 p=ROOT/'reviews/v03/external_bottle_rates.tsv'
 if not p.exists():
  sys.path.insert(0,str(ROOT/'reproducibility/v03'))
  from external_moisture import analyse
  analyse(OUT/'reji_development');p=OUT/'reji_development/external_bottle_rates.tsv'
 d=pd.read_csv(p,sep='\t')
 out=evaluate_temperate_folds(d);out.to_csv(OUT/'temperate_leave_soil_out.tsv',sep='\t',index=False)
 return [{'soil':soil,'model':model,'rmse':float(np.sqrt(np.mean((g.observed_k_min-g.predicted_k_min)**2)))} for (soil,model),g in out.groupby(['heldout_soil','model'])]

def run():
 OUT.mkdir(parents=True,exist_ok=True)
 report={'protocol':'reviews/v05/external-evaluation-protocol.md','seed':SEED,'sonoran':sonoran(),'namib':namib(),'temperate_leave_soil_out':temperate()}
 (OUT/'external_transfer.json').write_text(json.dumps(report,indent=2,allow_nan=False)+'\n')
 print(json.dumps({'sonoran':report['sonoran']['field_metrics'],'namib':report['namib']['contrasts'][:2],'temperate':report['temperate_leave_soil_out']},indent=2))
if __name__=='__main__':run()
